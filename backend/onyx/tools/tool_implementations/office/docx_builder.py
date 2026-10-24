"""Input model and builder for `create_document` (Word .docx)."""

from io import BytesIO
from typing import Annotated, Literal

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor
from docx.table import _Cell
from docx.text.paragraph import Paragraph
from pydantic import BaseModel, Field, JsonValue, model_validator

from onyx.tools.tool_implementations.office.branding import (
    OFFICE_BRAND,
    OfficeBrand,
)
from onyx.tools.tool_implementations.office.models import (
    MAX_CELL_TEXT_LENGTH,
    MAX_TEXT_LENGTH,
    MAX_TITLE_LENGTH,
    OfficeFileSpec,
    clean_text,
    parse_inline_runs,
)

DOCX_EXTENSION = ".docx"
MAX_DOCUMENT_BLOCKS = 300
MAX_LIST_ITEMS = 100
MAX_TABLE_COLUMNS = 20
MAX_TABLE_ROWS = 200
MAX_DOCUMENT_CHARS = 400_000
_TABLE_STYLE = "Table Grid"
_BRANDED_STYLES = ("Title", "Heading 1", "Heading 2", "Heading 3")


class HeadingBlock(BaseModel):
    type: Literal["heading"]
    text: str = Field(min_length=1, max_length=MAX_TITLE_LENGTH)
    level: int = Field(default=1, ge=1, le=3)


class ParagraphBlock(BaseModel):
    type: Literal["paragraph"]
    text: str = Field(min_length=1, max_length=MAX_TEXT_LENGTH)


class BulletListBlock(BaseModel):
    type: Literal["bullet_list"]
    items: list[Annotated[str, Field(max_length=MAX_TEXT_LENGTH)]] = Field(
        min_length=1, max_length=MAX_LIST_ITEMS
    )


class NumberedListBlock(BaseModel):
    type: Literal["numbered_list"]
    items: list[Annotated[str, Field(max_length=MAX_TEXT_LENGTH)]] = Field(
        min_length=1, max_length=MAX_LIST_ITEMS
    )


class TableBlock(BaseModel):
    type: Literal["table"]
    header: list[Annotated[str, Field(max_length=MAX_CELL_TEXT_LENGTH)]] = Field(
        min_length=1, max_length=MAX_TABLE_COLUMNS
    )
    rows: list[list[Annotated[str, Field(max_length=MAX_CELL_TEXT_LENGTH)]]] = Field(
        default_factory=list, max_length=MAX_TABLE_ROWS
    )

    @model_validator(mode="after")
    def _rows_match_header(self) -> "TableBlock":
        for index, row in enumerate(self.rows, start=1):
            if len(row) != len(self.header):
                raise ValueError(
                    f"table row {index} has {len(row)} cells; the header has "
                    f"{len(self.header)}"
                )
        return self


class PageBreakBlock(BaseModel):
    type: Literal["page_break"]


DocumentBlock = Annotated[
    HeadingBlock
    | ParagraphBlock
    | BulletListBlock
    | NumberedListBlock
    | TableBlock
    | PageBreakBlock,
    Field(discriminator="type"),
]


class DocumentSpec(OfficeFileSpec):
    title: str = Field(min_length=1, max_length=MAX_TITLE_LENGTH)
    subtitle: str | None = Field(default=None, max_length=MAX_TITLE_LENGTH)
    blocks: list[DocumentBlock] = Field(min_length=1, max_length=MAX_DOCUMENT_BLOCKS)

    @model_validator(mode="after")
    def _bounded_total_text(self) -> "DocumentSpec":
        total = sum(_block_chars(block) for block in self.blocks)
        if total > MAX_DOCUMENT_CHARS:
            raise ValueError(
                f"the document has {total} characters; the limit is "
                f"{MAX_DOCUMENT_CHARS}. Split it into several documents"
            )
        return self


def _block_chars(block: DocumentBlock) -> int:
    if isinstance(block, HeadingBlock | ParagraphBlock):
        return len(block.text)
    if isinstance(block, BulletListBlock | NumberedListBlock):
        return sum(len(item) for item in block.items)
    if isinstance(block, TableBlock):
        cells = [*block.header, *(cell for row in block.rows for cell in row)]
        return sum(len(cell) for cell in cells)
    return 0


def _add_runs(paragraph: Paragraph, text: str) -> None:
    for run in parse_inline_runs(text):
        styled = paragraph.add_run(run.text)
        styled.bold = run.bold or None
        styled.italic = run.italic or None


def _shade(cell: _Cell, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def _apply_brand(document: DocxDocument, brand: OfficeBrand) -> None:
    document.core_properties.author = brand.name
    color = RGBColor.from_string(brand.primary)
    for style_name in _BRANDED_STYLES:
        style = document.styles[style_name]
        style.font.color.rgb = color
        # The default Title style draws its bottom rule in the theme accent.
        for border in style.element.xpath(".//w:pBdr/w:bottom"):
            border.set(qn("w:color"), brand.primary)


def _add_table(
    document: DocxDocument, block: TableBlock, brand: OfficeBrand | None
) -> None:
    table = document.add_table(rows=1 + len(block.rows), cols=len(block.header))
    table.style = _TABLE_STYLE
    for col, header_text in enumerate(block.header):
        cell = table.cell(0, col)
        header_run = cell.paragraphs[0].add_run(clean_text(header_text))
        header_run.bold = True
        if brand is not None:
            _shade(cell, brand.primary)
            header_run.font.color.rgb = RGBColor.from_string(brand.on_primary)
    for row_index, row in enumerate(block.rows, start=1):
        for col, cell_text in enumerate(row):
            _add_runs(table.cell(row_index, col).paragraphs[0], cell_text)
    # Keep a paragraph between consecutive tables so Word does not merge them.
    document.add_paragraph()


def _add_block(
    document: DocxDocument, block: DocumentBlock, brand: OfficeBrand | None
) -> None:
    if isinstance(block, HeadingBlock):
        document.add_heading(clean_text(block.text), level=block.level)
    elif isinstance(block, ParagraphBlock):
        _add_runs(document.add_paragraph(), block.text)
    elif isinstance(block, BulletListBlock):
        for item in block.items:
            _add_runs(document.add_paragraph(style="List Bullet"), item)
    elif isinstance(block, NumberedListBlock):
        for item in block.items:
            _add_runs(document.add_paragraph(style="List Number"), item)
    elif isinstance(block, TableBlock):
        _add_table(document, block, brand)
    else:
        document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def build_docx(spec: DocumentSpec, brand: OfficeBrand | None = OFFICE_BRAND) -> bytes:
    document = Document()
    document.core_properties.title = clean_text(spec.title)
    if brand is not None:
        _apply_brand(document, brand)
    document.add_heading(clean_text(spec.title), level=0)
    if spec.subtitle:
        document.add_paragraph(clean_text(spec.subtitle), style="Subtitle")
    for block in spec.blocks:
        _add_block(document, block, brand)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def summarize_docx(spec: DocumentSpec) -> str:
    headings = sum(isinstance(block, HeadingBlock) for block in spec.blocks)
    tables = sum(isinstance(block, TableBlock) for block in spec.blocks)
    return (
        f"Word document '{spec.title}' with {len(spec.blocks)} blocks "
        f"({headings} headings, {tables} tables)."
    )


_TEXT_ITEMS: dict[str, JsonValue] = {
    "type": "array",
    "items": {"type": "string"},
    "maxItems": MAX_LIST_ITEMS,
}


def document_parameters(extension: str) -> dict[str, JsonValue]:
    """The document schema; the PDF tool shares it with its own extension."""
    return {
        "type": "object",
        "properties": {
            "file_name": {
                "type": "string",
                "description": "File name without a path, e.g. 'Q3 Report'. "
                f"The {extension} extension is added.",
            },
            **_DOCUMENT_BODY_PROPERTIES,
        },
        "required": ["file_name", "title", "blocks"],
    }


_DOCUMENT_BODY_PROPERTIES: dict[str, JsonValue] = {
    "title": {"type": "string", "description": "Document title."},
    "subtitle": {"type": "string", "description": "Optional subtitle."},
    "blocks": {
        "type": "array",
        "maxItems": MAX_DOCUMENT_BLOCKS,
        "description": "Document body in reading order. Each block sets "
        "'type' and only the fields for that type: heading (text, level "
        "1-3), paragraph (text; **bold** and *italic* allowed), "
        "bullet_list (items), numbered_list (items), table (header, rows; "
        "every row has as many cells as the header), page_break.",
        "items": {
            "type": "object",
            "properties": {
                "type": {
                    "type": "string",
                    "enum": [
                        "heading",
                        "paragraph",
                        "bullet_list",
                        "numbered_list",
                        "table",
                        "page_break",
                    ],
                },
                "text": {"type": "string"},
                "level": {"type": "integer", "minimum": 1, "maximum": 3},
                "items": _TEXT_ITEMS,
                "header": {
                    "type": "array",
                    "items": {"type": "string"},
                    "maxItems": MAX_TABLE_COLUMNS,
                },
                "rows": {
                    "type": "array",
                    "maxItems": MAX_TABLE_ROWS,
                    "items": {"type": "array", "items": {"type": "string"}},
                },
            },
            "required": ["type"],
        },
    },
}

DOCUMENT_PARAMETERS: dict[str, JsonValue] = document_parameters(DOCX_EXTENSION)
