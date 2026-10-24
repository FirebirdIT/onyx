"""Input model and builder for `create_presentation` (PowerPoint .pptx)."""

from io import BytesIO
from typing import Annotated, Literal

from pptx import Presentation
from pptx.chart.chart import Chart
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.presentation import Presentation as PptxPresentation
from pptx.slide import Slide
from pptx.text.text import TextFrame
from pptx.util import Emu, Inches, Pt
from pydantic import BaseModel, Field, FiniteFloat, JsonValue, model_validator

from onyx.tools.tool_implementations.office.branding import (
    OFFICE_BRAND,
    OfficeBrand,
)
from onyx.tools.tool_implementations.office.models import (
    MAX_NOTES_LENGTH,
    MAX_TITLE_LENGTH,
    OfficeFileSpec,
    clean_text,
    parse_inline_runs,
)

PPTX_EXTENSION = ".pptx"
MAX_SLIDES = 50
MAX_SLIDE_TITLE_LENGTH = 200
MAX_BULLETS = 10
MAX_BULLET_LENGTH = 300
MAX_TABLE_COLUMNS = 8
MAX_TABLE_ROWS = 12
MAX_TABLE_CELL_LENGTH = 200
MAX_CHART_CATEGORIES = 50
MAX_CHART_SERIES = 8

# Widescreen 16:9, 13.333 x 7.5 inches.
_SLIDE_WIDTH = Emu(12_192_000)
_SLIDE_HEIGHT = Emu(6_858_000)
# Layout indexes of the python-pptx default template.
_LAYOUT_TITLE = 0
_LAYOUT_TITLE_AND_CONTENT = 1
_LAYOUT_SECTION = 2
_LAYOUT_COMPARISON = 4
_LAYOUT_TITLE_ONLY = 5
# Placeholder boxes as (left, top, width, height) in inches.
_Box = tuple[float, float, float, float]
_TITLE_BOX: _Box = (0.6, 0.35, 12.13, 1.1)
_BODY_BOX: _Box = (0.6, 1.6, 12.13, 5.4)
_GEOMETRY: dict[int, dict[int, _Box]] = {
    _LAYOUT_TITLE: {0: (0.9, 2.1, 11.53, 1.6), 1: (0.9, 3.9, 11.53, 1.2)},
    _LAYOUT_SECTION: {0: (0.9, 2.5, 11.53, 1.4), 1: (0.9, 4.1, 11.53, 1.0)},
    _LAYOUT_TITLE_AND_CONTENT: {0: _TITLE_BOX, 1: _BODY_BOX},
    _LAYOUT_COMPARISON: {
        0: _TITLE_BOX,
        1: (0.6, 1.6, 5.9, 0.6),
        2: (0.6, 2.3, 5.9, 4.7),
        3: (6.83, 1.6, 5.9, 0.6),
        4: (6.83, 2.3, 5.9, 4.7),
    },
    _LAYOUT_TITLE_ONLY: {0: _TITLE_BOX},
}
_TABLE_ROW_HEIGHT_IN = 0.45
_TABLE_FONT_SIZE = Pt(14)
# Brand band across the top edge of every slide.
_BAND_HEIGHT = Inches(0.12)

Bullet = Annotated[str, Field(min_length=1, max_length=MAX_BULLET_LENGTH)]
Cell = Annotated[str, Field(max_length=MAX_TABLE_CELL_LENGTH)]
SlideTitle = Annotated[str, Field(min_length=1, max_length=MAX_SLIDE_TITLE_LENGTH)]


class _SlideBase(BaseModel):
    title: SlideTitle
    notes: str | None = Field(default=None, max_length=MAX_NOTES_LENGTH)


class TitleSlide(_SlideBase):
    layout: Literal["title"]
    subtitle: str | None = Field(default=None, max_length=MAX_TITLE_LENGTH)


class SectionSlide(_SlideBase):
    layout: Literal["section"]
    subtitle: str | None = Field(default=None, max_length=MAX_TITLE_LENGTH)


class BulletsSlide(_SlideBase):
    layout: Literal["bullets"]
    bullets: list[Bullet] = Field(min_length=1, max_length=MAX_BULLETS)


class TwoColumnSlide(_SlideBase):
    layout: Literal["two_column"]
    left_heading: str | None = Field(default=None, max_length=MAX_SLIDE_TITLE_LENGTH)
    left_bullets: list[Bullet] = Field(min_length=1, max_length=MAX_BULLETS)
    right_heading: str | None = Field(default=None, max_length=MAX_SLIDE_TITLE_LENGTH)
    right_bullets: list[Bullet] = Field(min_length=1, max_length=MAX_BULLETS)


class TableSlide(_SlideBase):
    layout: Literal["table"]
    header: list[Cell] = Field(min_length=1, max_length=MAX_TABLE_COLUMNS)
    rows: list[list[Cell]] = Field(min_length=1, max_length=MAX_TABLE_ROWS)

    @model_validator(mode="after")
    def _rows_match_header(self) -> "TableSlide":
        for index, row in enumerate(self.rows, start=1):
            if len(row) != len(self.header):
                raise ValueError(
                    f"table row {index} has {len(row)} cells; the header has "
                    f"{len(self.header)}"
                )
        return self


class ChartSeries(BaseModel):
    name: Annotated[str, Field(min_length=1, max_length=MAX_SLIDE_TITLE_LENGTH)]
    values: list[FiniteFloat] = Field(min_length=1, max_length=MAX_CHART_CATEGORIES)


class ChartSlide(_SlideBase):
    layout: Literal["chart"]
    chart_type: Literal["bar", "line", "pie"]
    categories: list[Cell] = Field(min_length=1, max_length=MAX_CHART_CATEGORIES)
    series: list[ChartSeries] = Field(min_length=1, max_length=MAX_CHART_SERIES)

    @model_validator(mode="after")
    def _series_match_categories(self) -> "ChartSlide":
        for series in self.series:
            if len(series.values) != len(self.categories):
                raise ValueError(
                    f"series '{series.name}' has {len(series.values)} values; "
                    f"there are {len(self.categories)} categories"
                )
        if self.chart_type == "pie":
            values = self.series[0].values
            if len(self.series) != 1 or min(values) < 0 or sum(values) <= 0:
                raise ValueError(
                    "a pie chart takes one series of non-negative values with a "
                    "positive total"
                )
        return self


SlideSpec = Annotated[
    TitleSlide | SectionSlide | BulletsSlide | TwoColumnSlide | TableSlide | ChartSlide,
    Field(discriminator="layout"),
]


class PresentationSpec(OfficeFileSpec):
    title: str = Field(min_length=1, max_length=MAX_TITLE_LENGTH)
    slides: list[SlideSpec] = Field(min_length=1, max_length=MAX_SLIDES)


def _place(slide: Slide, layout_index: int) -> None:
    """Fit the template's 4:3 placeholders to the 16:9 slide."""
    boxes = _GEOMETRY[layout_index]
    for placeholder in slide.placeholders:
        box = boxes.get(placeholder.placeholder_format.idx)
        if box is None:
            continue
        left, top, width, height = box
        placeholder.left = Inches(left)
        placeholder.top = Inches(top)
        placeholder.width = Inches(width)
        placeholder.height = Inches(height)


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def _add_band(slide: Slide, brand: OfficeBrand) -> None:
    band = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Emu(0), Emu(0), _SLIDE_WIDTH, _BAND_HEIGHT
    )
    band.fill.solid()
    band.fill.fore_color.rgb = _rgb(brand.primary)
    band.line.fill.background()


def _new_slide(
    prs: PptxPresentation,
    layout_index: int,
    title: str,
    brand: OfficeBrand | None,
) -> Slide:
    slide = prs.slides.add_slide(prs.slide_layouts[layout_index])
    _place(slide, layout_index)
    title_shape = slide.shapes.title
    if title_shape is not None:
        title_shape.text_frame.text = clean_text(title)
        if brand is not None:
            for paragraph in title_shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.color.rgb = _rgb(brand.primary)
    if brand is not None:
        _add_band(slide, brand)
    return slide


def _fill_text(slide: Slide, idx: int, text: str | None) -> None:
    """Set a placeholder's text, or remove the placeholder when `text` is empty."""
    placeholder = slide.placeholders[idx]
    if text:
        placeholder.text_frame.text = clean_text(text)
        return
    element = placeholder.element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def _fill_bullets(frame: TextFrame, bullets: list[str]) -> None:
    frame.word_wrap = True
    for index, bullet in enumerate(bullets):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        for run in parse_inline_runs(bullet):
            styled = paragraph.add_run()
            styled.text = run.text
            styled.font.bold = run.bold or None
            styled.font.italic = run.italic or None


def _add_table(slide: Slide, spec: TableSlide, brand: OfficeBrand | None) -> None:
    left, top, width, _ = _BODY_BOX
    height = _TABLE_ROW_HEIGHT_IN * (len(spec.rows) + 1)
    shape = slide.shapes.add_table(
        len(spec.rows) + 1,
        len(spec.header),
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    table = shape.table
    for row_index, row in enumerate([spec.header, *spec.rows]):
        for col, text in enumerate(row):
            frame = table.cell(row_index, col).text_frame
            frame.text = clean_text(text)
            for paragraph in frame.paragraphs:
                paragraph.font.size = _TABLE_FONT_SIZE
            if brand is None:
                continue
            cell = table.cell(row_index, col)
            cell.fill.solid()
            if row_index == 0:
                cell.fill.fore_color.rgb = _rgb(brand.primary)
                for paragraph in frame.paragraphs:
                    paragraph.font.bold = True
                    paragraph.font.color.rgb = _rgb(brand.on_primary)
            else:
                band = brand.tint if row_index % 2 == 0 else "FFFFFF"
                cell.fill.fore_color.rgb = _rgb(band)


def _color_chart(chart: Chart, chart_type: str, brand: OfficeBrand) -> None:
    colors = brand.chart_colors
    plot = chart.plots[0]
    if chart_type == "pie":
        for index, point in enumerate(plot.series[0].points):
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = _rgb(colors[index % len(colors)])
        return
    for index, series in enumerate(plot.series):
        color = _rgb(colors[index % len(colors)])
        if chart_type == "line":
            series.format.line.color.rgb = color
            series.marker.format.fill.solid()
            series.marker.format.fill.fore_color.rgb = color
        else:
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = color


def _add_chart(slide: Slide, spec: ChartSlide, brand: OfficeBrand | None) -> None:
    data = CategoryChartData()
    data.categories = [clean_text(category) for category in spec.categories]
    for series in spec.series:
        data.add_series(clean_text(series.name), series.values)
    chart_type = {
        "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "line": XL_CHART_TYPE.LINE_MARKERS,
        "pie": XL_CHART_TYPE.PIE,
    }[spec.chart_type]
    left, top, width, height = _BODY_BOX
    chart = slide.shapes.add_chart(
        chart_type, Inches(left), Inches(top), Inches(width), Inches(height), data
    ).chart
    chart.has_legend = spec.chart_type == "pie" or len(spec.series) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    if brand is not None:
        _color_chart(chart, spec.chart_type, brand)


def _add_slide(
    prs: PptxPresentation, spec: SlideSpec, brand: OfficeBrand | None
) -> Slide:
    if isinstance(spec, TitleSlide | SectionSlide):
        layout = _LAYOUT_TITLE if isinstance(spec, TitleSlide) else _LAYOUT_SECTION
        slide = _new_slide(prs, layout, spec.title, brand)
        _fill_text(slide, 1, spec.subtitle)
    elif isinstance(spec, BulletsSlide):
        slide = _new_slide(prs, _LAYOUT_TITLE_AND_CONTENT, spec.title, brand)
        _fill_bullets(slide.placeholders[1].text_frame, spec.bullets)
    elif isinstance(spec, TwoColumnSlide):
        slide = _new_slide(prs, _LAYOUT_COMPARISON, spec.title, brand)
        _fill_bullets(slide.placeholders[2].text_frame, spec.left_bullets)
        _fill_bullets(slide.placeholders[4].text_frame, spec.right_bullets)
        _fill_text(slide, 1, spec.left_heading)
        _fill_text(slide, 3, spec.right_heading)
    elif isinstance(spec, TableSlide):
        slide = _new_slide(prs, _LAYOUT_TITLE_ONLY, spec.title, brand)
        _add_table(slide, spec, brand)
    else:
        slide = _new_slide(prs, _LAYOUT_TITLE_ONLY, spec.title, brand)
        _add_chart(slide, spec, brand)
    if spec.notes:
        slide.notes_slide.notes_text_frame.text = clean_text(spec.notes)
    return slide


def build_pptx(
    spec: PresentationSpec, brand: OfficeBrand | None = OFFICE_BRAND
) -> bytes:
    prs = Presentation()
    prs.slide_width = _SLIDE_WIDTH
    prs.slide_height = _SLIDE_HEIGHT
    prs.core_properties.title = clean_text(spec.title)
    if brand is not None:
        prs.core_properties.author = brand.name
    for slide_spec in spec.slides:
        _add_slide(prs, slide_spec, brand)
    buffer = BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def summarize_pptx(spec: PresentationSpec) -> str:
    with_notes = sum(1 for slide in spec.slides if slide.notes)
    return (
        f"PowerPoint presentation '{spec.title}' with {len(spec.slides)} slides "
        f"({with_notes} with speaker notes)."
    )


_STRINGS: dict[str, JsonValue] = {"type": "array", "items": {"type": "string"}}

PRESENTATION_PARAMETERS: dict[str, JsonValue] = {
    "type": "object",
    "properties": {
        "file_name": {
            "type": "string",
            "description": "File name without a path. The .pptx extension is added.",
        },
        "title": {"type": "string", "description": "Presentation title."},
        "slides": {
            "type": "array",
            "maxItems": MAX_SLIDES,
            "description": "Slides in order. Each slide sets 'layout', 'title', "
            "optional 'notes' (speaker notes) and only the fields for its layout: "
            "title (subtitle), section (subtitle), bullets (bullets), two_column "
            "(left_heading, left_bullets, right_heading, right_bullets), table "
            "(header, rows; every row has as many cells as the header), chart "
            "(chart_type, categories, series; each series has one value per "
            f"category; pie takes one series). At most {MAX_BULLETS} short "
            f"bullets and {MAX_TABLE_ROWS} table rows per slide; split longer "
            "content across slides.",
            "items": {
                "type": "object",
                "properties": {
                    "layout": {
                        "type": "string",
                        "enum": [
                            "title",
                            "section",
                            "bullets",
                            "two_column",
                            "table",
                            "chart",
                        ],
                    },
                    "title": {"type": "string"},
                    "subtitle": {"type": "string"},
                    "notes": {"type": "string"},
                    "bullets": _STRINGS,
                    "left_heading": {"type": "string"},
                    "left_bullets": _STRINGS,
                    "right_heading": {"type": "string"},
                    "right_bullets": _STRINGS,
                    "header": _STRINGS,
                    "rows": {"type": "array", "items": _STRINGS},
                    "chart_type": {"type": "string", "enum": ["bar", "line", "pie"]},
                    "categories": _STRINGS,
                    "series": {
                        "type": "array",
                        "items": {
                            "type": "object",
                            "properties": {
                                "name": {"type": "string"},
                                "values": {
                                    "type": "array",
                                    "items": {"type": "number"},
                                },
                            },
                            "required": ["name", "values"],
                        },
                    },
                },
                "required": ["layout", "title"],
            },
        },
    },
    "required": ["file_name", "title", "slides"],
}
