"""Input model and builder for `create_spreadsheet` (Excel .xlsx)."""

import re
from io import BytesIO
from typing import Annotated, Literal

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_from_string, range_boundaries
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import (
    BaseModel,
    Field,
    FiniteFloat,
    JsonValue,
    StrictBool,
    StrictInt,
    field_validator,
    model_validator,
)

from onyx.tools.tool_implementations.office.models import (
    MAX_CELL_TEXT_LENGTH,
    MAX_TITLE_LENGTH,
    OfficeFileSpec,
    clean_text,
)

XLSX_EXTENSION = ".xlsx"
MAX_SHEETS = 20
MAX_COLUMNS = 50
MAX_ROWS_PER_SHEET = 5_000
MAX_WORKBOOK_CELLS = 200_000
MAX_CHARTS_PER_SHEET = 5
MAX_SHEET_NAME_LENGTH = 31
MAX_COLUMN_WIDTH = 255
# Excel's own limit for formula text.
MAX_FORMULA_LENGTH = 8_192
_CHART_HEIGHT_ROWS = 16
_HEADER_FILL = PatternFill(start_color="DDE4EE", end_color="DDE4EE", fill_type="solid")
_INVALID_SHEET_NAME_CHARS = re.compile(r"[\[\]:\*\?/\\]")
# Dynamic Data Exchange formulas (=cmd|'...'!A0) can start programs.
_DDE_FORMULA = re.compile(r"^=\s*[-+@]*\s*[A-Za-z0-9_.]+\s*\|")
_CELL_RANGE = re.compile(r"^\$?[A-Z]{1,3}\$?\d+:\$?[A-Z]{1,3}\$?\d+$")

CellValue = (
    StrictBool
    | StrictInt
    | FiniteFloat
    | Annotated[str, Field(max_length=MAX_FORMULA_LENGTH)]
    | None
)


class ChartSpec(BaseModel):
    type: Literal["bar", "line", "pie"]
    title: str | None = Field(default=None, max_length=MAX_TITLE_LENGTH)
    data_range: str
    categories_range: str | None = None
    anchor: str | None = None

    @field_validator("data_range", "categories_range")
    @classmethod
    def _is_a1_range(cls, value: str | None) -> str | None:
        if value is not None and not _CELL_RANGE.match(value.upper()):
            raise ValueError(f"'{value}' is not an A1 range such as B1:C10")
        return value.upper() if value else value

    @field_validator("anchor")
    @classmethod
    def _is_cell(cls, value: str | None) -> str | None:
        if value is None:
            return None
        try:
            coordinate_from_string(value.upper())
        except ValueError:
            raise ValueError(f"'{value}' is not a cell such as H2")
        return value.upper()


class SheetSpec(BaseModel):
    name: str = Field(min_length=1, max_length=MAX_SHEET_NAME_LENGTH)
    header: list[Annotated[str, Field(max_length=MAX_CELL_TEXT_LENGTH)]] = Field(
        min_length=1, max_length=MAX_COLUMNS
    )
    rows: list[list[CellValue]] = Field(
        default_factory=list, max_length=MAX_ROWS_PER_SHEET
    )
    column_widths: list[Annotated[float, Field(ge=1, le=MAX_COLUMN_WIDTH)]] | None = (
        Field(default=None, max_length=MAX_COLUMNS)
    )
    freeze_header: bool = True
    charts: list[ChartSpec] = Field(
        default_factory=list, max_length=MAX_CHARTS_PER_SHEET
    )

    @field_validator("name")
    @classmethod
    def _valid_sheet_name(cls, value: str) -> str:
        if _INVALID_SHEET_NAME_CHARS.search(value) or value.startswith("'"):
            raise ValueError(
                f"sheet name '{value}' may not contain [ ] : * ? / \\ or start "
                "with an apostrophe"
            )
        return value

    @model_validator(mode="after")
    def _rows_and_charts_fit(self) -> "SheetSpec":
        width = len(self.header)
        for index, row in enumerate(self.rows, start=1):
            if len(row) > width:
                raise ValueError(
                    f"sheet '{self.name}' row {index} has {len(row)} cells; the "
                    f"header has {width}"
                )
            for value in row:
                _check_cell(value)
        if self.column_widths and len(self.column_widths) > width:
            raise ValueError(f"sheet '{self.name}' has more column widths than columns")
        for chart in self.charts:
            _check_chart(chart, self.name, width, len(self.rows) + 1)
        return self


def _check_cell(value: CellValue) -> None:
    if not isinstance(value, str):
        return
    if value.startswith("=") and _DDE_FORMULA.match(value):
        raise ValueError("DDE formulas (for example =cmd|...) are not allowed")
    if not value.startswith("=") and len(value) > MAX_CELL_TEXT_LENGTH:
        raise ValueError(
            f"a text cell has {len(value)} characters; the limit is "
            f"{MAX_CELL_TEXT_LENGTH}"
        )


def _check_chart(chart: ChartSpec, sheet: str, max_col: int, max_row: int) -> None:
    ranges = [chart.data_range]
    if chart.categories_range:
        ranges.append(chart.categories_range)
    for cell_range in ranges:
        min_c, min_r, max_c, max_r = range_boundaries(cell_range)
        if (
            min_c is None
            or min_r is None
            or max_c is None
            or max_r is None
            or max_c > max_col
            or max_r > max_row
        ):
            raise ValueError(
                f"chart range {cell_range} is outside the data of sheet '{sheet}'"
            )
    min_c, _, max_c, _ = range_boundaries(chart.data_range)
    if chart.type == "pie" and min_c != max_c:
        raise ValueError("a pie chart takes one data column")


class WorkbookSpec(OfficeFileSpec):
    sheets: list[SheetSpec] = Field(min_length=1, max_length=MAX_SHEETS)

    @model_validator(mode="after")
    def _unique_names_and_bounded_size(self) -> "WorkbookSpec":
        names = [sheet.name.lower() for sheet in self.sheets]
        if len(set(names)) != len(names):
            raise ValueError("sheet names must be unique (ignoring case)")
        cells = sum(len(s.header) * (len(s.rows) + 1) for s in self.sheets)
        if cells > MAX_WORKBOOK_CELLS:
            raise ValueError(
                f"the workbook has {cells} cells; the limit is {MAX_WORKBOOK_CELLS}"
            )
        return self


def _cell(value: CellValue) -> str | int | float | bool | None:
    return clean_text(value) if isinstance(value, str) else value


def _add_chart(sheet: Worksheet, spec: ChartSpec, anchor: str) -> None:
    chart: BarChart | LineChart | PieChart
    if spec.type == "bar":
        chart = BarChart()
    elif spec.type == "line":
        chart = LineChart()
    else:
        chart = PieChart()
    if spec.title:
        chart.title = clean_text(spec.title)
    min_c, min_r, max_c, max_r = range_boundaries(spec.data_range)
    chart.add_data(
        Reference(sheet, min_col=min_c, min_row=min_r, max_col=max_c, max_row=max_r),
        titles_from_data=True,
    )
    if spec.categories_range:
        cat_min_c, cat_min_r, cat_max_c, cat_max_r = range_boundaries(
            spec.categories_range
        )
        chart.set_categories(
            Reference(
                sheet,
                min_col=cat_min_c,
                min_row=cat_min_r,
                max_col=cat_max_c,
                max_row=cat_max_r,
            )
        )
    sheet.add_chart(chart, anchor)


def _fill_sheet(sheet: Worksheet, spec: SheetSpec) -> None:
    sheet.append([clean_text(text) for text in spec.header])
    for column in range(1, len(spec.header) + 1):
        header_cell = sheet.cell(row=1, column=column)
        header_cell.font = Font(bold=True)
        header_cell.fill = _HEADER_FILL
    for row in spec.rows:
        sheet.append([_cell(value) for value in row])
    for index, width in enumerate(spec.column_widths or [], start=1):
        # The openpyxl stubs omit Worksheet.column_dimensions.
        dimensions = sheet.column_dimensions  # ty: ignore[unresolved-attribute]
        dimensions[get_column_letter(index)].width = width
    if spec.freeze_header:
        sheet.freeze_panes = "A2"
    default_column = get_column_letter(len(spec.header) + 2)
    for index, chart in enumerate(spec.charts):
        anchor = chart.anchor or f"{default_column}{2 + index * _CHART_HEIGHT_ROWS}"
        _add_chart(sheet, chart, anchor)


def build_xlsx(spec: WorkbookSpec) -> bytes:
    workbook = Workbook()
    first = workbook.active
    if isinstance(first, Worksheet):
        workbook.remove(first)
    for sheet_spec in spec.sheets:
        _fill_sheet(workbook.create_sheet(title=sheet_spec.name), sheet_spec)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def summarize_xlsx(spec: WorkbookSpec) -> str:
    parts = [
        f"'{sheet.name}' ({len(sheet.rows)} rows, {len(sheet.charts)} charts)"
        for sheet in spec.sheets
    ]
    return f"Excel workbook with {len(spec.sheets)} sheets: {', '.join(parts)}."


_CELL_SCHEMA: dict[str, JsonValue] = {
    "anyOf": [
        {"type": "string"},
        {"type": "number"},
        {"type": "boolean"},
        {"type": "null"},
    ]
}

SPREADSHEET_PARAMETERS: dict[str, JsonValue] = {
    "type": "object",
    "properties": {
        "file_name": {
            "type": "string",
            "description": "File name without a path. The .xlsx extension is added.",
        },
        "sheets": {
            "type": "array",
            "maxItems": MAX_SHEETS,
            "items": {
                "type": "object",
                "properties": {
                    "name": {
                        "type": "string",
                        "description": "Sheet name, at most 31 characters, "
                        "unique, without [ ] : * ? / \\",
                    },
                    "header": {
                        "type": "array",
                        "items": {"type": "string"},
                        "maxItems": MAX_COLUMNS,
                        "description": "Column titles. They fill row 1.",
                    },
                    "rows": {
                        "type": "array",
                        "maxItems": MAX_ROWS_PER_SHEET,
                        "description": "Data rows, starting at row 2. Cells are "
                        "strings, numbers, booleans or null. A string that starts "
                        "with '=' is stored as an Excel formula, e.g. '=SUM(B2:B5)'. "
                        "Use real numbers, not formatted strings, for numeric data.",
                        "items": {"type": "array", "items": _CELL_SCHEMA},
                    },
                    "column_widths": {
                        "type": "array",
                        "items": {"type": "number"},
                        "description": "Optional widths in characters, by column.",
                    },
                    "freeze_header": {
                        "type": "boolean",
                        "description": "Keep row 1 visible when scrolling. "
                        "Default true.",
                    },
                    "charts": {
                        "type": "array",
                        "maxItems": MAX_CHARTS_PER_SHEET,
                        "items": {
                            "type": "object",
                            "properties": {
                                "type": {
                                    "type": "string",
                                    "enum": ["bar", "line", "pie"],
                                },
                                "title": {"type": "string"},
                                "data_range": {
                                    "type": "string",
                                    "description": "A1 range of the series, "
                                    "including the header row for series names, "
                                    "e.g. 'B1:C6'. A pie chart takes one column.",
                                },
                                "categories_range": {
                                    "type": "string",
                                    "description": "A1 range of the category "
                                    "labels without the header, e.g. 'A2:A6'.",
                                },
                                "anchor": {
                                    "type": "string",
                                    "description": "Optional top-left cell, e.g. "
                                    "'H2'. Default is beside the data.",
                                },
                            },
                            "required": ["type", "data_range"],
                        },
                    },
                },
                "required": ["name", "header", "rows"],
            },
        },
    },
    "required": ["file_name", "sheets"],
}
