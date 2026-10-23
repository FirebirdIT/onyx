from io import BytesIO
from typing import Any

import pytest
from pptx import Presentation
from pydantic import ValidationError

from onyx.tools.tool_implementations.office.pptx_builder import (
    MAX_BULLETS,
    MAX_SLIDES,
    PresentationSpec,
    build_pptx,
)


def _slides() -> list[dict[str, Any]]:
    return [
        {
            "layout": "title",
            "title": "Strategy 2027",
            "subtitle": "Board review",
            "notes": "Welcome everyone.",
        },
        {"layout": "section", "title": "Market"},
        {
            "layout": "bullets",
            "title": "Highlights",
            "bullets": ["Revenue **up**", "Costs down"],
        },
        {
            "layout": "two_column",
            "title": "Options",
            "left_heading": "Build",
            "left_bullets": ["Control"],
            "right_bullets": ["Speed"],
        },
        {
            "layout": "table",
            "title": "Budget",
            "header": ["Item", "Cost"],
            "rows": [["Hiring", "10"], ["Tools", "2"]],
        },
        {
            "layout": "chart",
            "title": "Growth",
            "chart_type": "line",
            "categories": ["2025", "2026"],
            "series": [
                {"name": "Revenue", "values": [1, 2]},
                {"name": "Cost", "values": [1, 1.5]},
            ],
            "notes": "Show the trend.",
        },
    ]


def _spec(slides: list[dict[str, Any]]) -> PresentationSpec:
    return PresentationSpec.model_validate(
        {"file_name": "deck", "title": "Strategy", "slides": slides}
    )


def _slide_texts(slide: Any) -> list[str]:
    return [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]


def test_build_pptx_round_trip() -> None:
    prs = Presentation(BytesIO(build_pptx(_spec(_slides()))))

    assert (prs.slide_width, prs.slide_height) == (12_192_000, 6_858_000)
    slides = list(prs.slides)
    assert len(slides) == 6
    titles = [slide.shapes.title.text for slide in slides]
    assert titles == [
        "Strategy 2027",
        "Market",
        "Highlights",
        "Options",
        "Budget",
        "Growth",
    ]
    assert "Board review" in _slide_texts(slides[0])
    assert slides[0].notes_slide.notes_text_frame.text == "Welcome everyone."
    assert slides[5].notes_slide.notes_text_frame.text == "Show the trend."
    assert not slides[1].has_notes_slide

    assert "Revenue up\nCosts down" in _slide_texts(slides[2])
    two_column = _slide_texts(slides[3])
    assert "Build" in two_column
    assert "Control" in two_column
    assert "Speed" in two_column

    table = next(s for s in slides[4].shapes if s.has_table).table
    assert table.cell(0, 1).text == "Cost"
    assert table.cell(2, 0).text == "Tools"

    chart = next(s for s in slides[5].shapes if s.has_chart).chart
    assert [series.name for series in chart.plots[0].series] == ["Revenue", "Cost"]
    assert list(chart.plots[0].categories) == ["2025", "2026"]


@pytest.mark.parametrize(
    "slide",
    [
        {"layout": "bullets", "title": "T", "bullets": ["x"] * (MAX_BULLETS + 1)},
        {"layout": "bullets", "title": "T", "bullets": []},
        {"layout": "table", "title": "T", "header": ["a", "b"], "rows": [["1"]]},
        {
            "layout": "chart",
            "title": "T",
            "chart_type": "bar",
            "categories": ["a", "b"],
            "series": [{"name": "s", "values": [1]}],
        },
        {
            "layout": "chart",
            "title": "T",
            "chart_type": "pie",
            "categories": ["a", "b"],
            "series": [{"name": "s", "values": [-1, 2]}],
        },
        {"layout": "video", "title": "T"},
        {"layout": "title", "title": ""},
    ],
)
def test_rejects_invalid_slide(slide: dict[str, Any]) -> None:
    with pytest.raises(ValidationError):
        _spec([slide])


def test_rejects_too_many_slides() -> None:
    with pytest.raises(ValidationError):
        _spec([{"layout": "section", "title": "S"}] * (MAX_SLIDES + 1))
