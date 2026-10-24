"""The brand colours reach each office format; no brand keeps the defaults."""

from io import BytesIO
from typing import Any

from docx import Document
from docx.oxml.ns import qn
from openpyxl import load_workbook
from pptx import Presentation

from onyx.tools.tool_implementations.office.branding import DEEPX_OFFICE_BRAND
from onyx.tools.tool_implementations.office.docx_builder import (
    DocumentSpec,
    build_docx,
)
from onyx.tools.tool_implementations.office.pptx_builder import (
    PresentationSpec,
    build_pptx,
)
from onyx.tools.tool_implementations.office.xlsx_builder import (
    WorkbookSpec,
    build_xlsx,
)

BRAND = DEEPX_OFFICE_BRAND


def _document() -> DocumentSpec:
    return DocumentSpec.model_validate(
        {
            "file_name": "Report",
            "title": "Report",
            "blocks": [
                {"type": "heading", "text": "Numbers", "level": 1},
                {"type": "table", "header": ["A", "B"], "rows": [["1", "2"]]},
            ],
        }
    )


def _presentation() -> PresentationSpec:
    slides: list[dict[str, Any]] = [
        {"layout": "title", "title": "Deck"},
        {
            "layout": "table",
            "title": "Table",
            "header": ["A", "B"],
            "rows": [["1", "2"], ["3", "4"]],
        },
        {
            "layout": "chart",
            "title": "Chart",
            "chart_type": "bar",
            "categories": ["Q1", "Q2"],
            "series": [{"name": "Sales", "values": [1, 2]}],
        },
        {
            "layout": "chart",
            "title": "Pie",
            "chart_type": "pie",
            "categories": ["A", "B"],
            "series": [{"name": "Share", "values": [60, 40]}],
        },
        {
            "layout": "chart",
            "title": "Line",
            "chart_type": "line",
            "categories": ["Q1", "Q2"],
            "series": [{"name": "Trend", "values": [1, 3]}],
        },
    ]
    return PresentationSpec.model_validate(
        {"file_name": "Deck", "title": "Deck", "slides": slides}
    )


def _workbook() -> WorkbookSpec:
    return WorkbookSpec.model_validate(
        {
            "file_name": "Book",
            "sheets": [{"name": "S", "header": ["A"], "rows": [[1]]}],
        }
    )


def _header_shading(content: bytes) -> str | None:
    cell = Document(BytesIO(content)).tables[0].cell(0, 0)
    shading = cell._tc.tcPr.find(qn("w:shd")) if cell._tc.tcPr is not None else None
    return None if shading is None else shading.get(qn("w:fill"))


def test_docx_brand() -> None:
    content = build_docx(_document(), brand=BRAND)
    document = Document(BytesIO(content))
    assert document.core_properties.author == "Deep X"
    assert str(document.styles["Heading 1"].font.color.rgb) == BRAND.primary
    assert _header_shading(content) == BRAND.primary


def test_docx_without_brand_keeps_defaults() -> None:
    content = build_docx(_document(), brand=None)
    document = Document(BytesIO(content))
    assert document.styles["Heading 1"].font.color.rgb != BRAND.primary
    assert _header_shading(content) is None


def test_pptx_brand() -> None:
    prs = Presentation(BytesIO(build_pptx(_presentation(), brand=BRAND)))
    assert prs.core_properties.author == "Deep X"
    title_slide, table_slide = prs.slides[0], prs.slides[1]
    title_run = title_slide.shapes.title.text_frame.paragraphs[0].runs[0]
    assert str(title_run.font.color.rgb) == BRAND.primary
    # Every slide carries the brand band.
    for slide in prs.slides:
        assert any(
            shape.shape_type is not None
            and shape.top == 0
            and shape.has_text_frame
            and not shape.is_placeholder
            for shape in slide.shapes
        )
    table = next(s for s in table_slide.shapes if s.has_table).table
    assert str(table.cell(0, 0).fill.fore_color.rgb) == BRAND.primary
    assert str(table.cell(2, 0).fill.fore_color.rgb) == BRAND.tint


def test_pptx_without_brand_has_no_band() -> None:
    prs = Presentation(BytesIO(build_pptx(_presentation(), brand=None)))
    for slide in prs.slides:
        assert all(shape.is_placeholder or shape.top != 0 for shape in slide.shapes)


def test_xlsx_brand() -> None:
    sheet = load_workbook(BytesIO(build_xlsx(_workbook(), brand=BRAND)))["S"]
    header = sheet.cell(row=1, column=1)
    assert header.fill.start_color.rgb.endswith(BRAND.primary)
    assert header.font.color.rgb.endswith(BRAND.on_primary)


def test_xlsx_without_brand_keeps_defaults() -> None:
    sheet = load_workbook(BytesIO(build_xlsx(_workbook(), brand=None)))["S"]
    assert sheet.cell(row=1, column=1).fill.start_color.rgb.endswith("DDE4EE")
