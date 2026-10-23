from io import BytesIO
from typing import Any

import pytest
from docx import Document
from pydantic import ValidationError

from onyx.tools.tool_implementations.office.docx_builder import (
    MAX_DOCUMENT_BLOCKS,
    DocumentSpec,
    build_docx,
)
from onyx.tools.tool_implementations.office.models import (
    parse_inline_runs,
    sanitize_file_name,
)


def _spec(**overrides: Any) -> DocumentSpec:
    payload: dict[str, Any] = {
        "file_name": "Quarterly Report",
        "title": "Quarterly Report",
        "subtitle": "Q3 2026",
        "blocks": [
            {"type": "heading", "text": "Summary", "level": 1},
            {"type": "paragraph", "text": "Revenue grew **12%** in *Q3*."},
            {"type": "bullet_list", "items": ["New region", "New product"]},
            {"type": "numbered_list", "items": ["Hire", "Ship"]},
            {"type": "page_break"},
            {"type": "heading", "text": "Numbers", "level": 2},
            {
                "type": "table",
                "header": ["Region", "Revenue"],
                "rows": [["EMEA", "10"], ["APAC", "12"]],
            },
        ],
    }
    payload.update(overrides)
    return DocumentSpec.model_validate(payload)


def test_build_docx_round_trip() -> None:
    document = Document(BytesIO(build_docx(_spec())))

    paragraphs = [
        (p.style.name if p.style else None, p.text)
        for p in document.paragraphs
        if p.text
    ]
    assert ("Title", "Quarterly Report") in paragraphs
    assert ("Subtitle", "Q3 2026") in paragraphs
    assert ("Heading 1", "Summary") in paragraphs
    assert ("Heading 2", "Numbers") in paragraphs
    assert ("List Bullet", "New region") in paragraphs
    assert ("List Number", "Ship") in paragraphs
    assert document.core_properties.title == "Quarterly Report"

    body = next(p for p in document.paragraphs if p.text.startswith("Revenue"))
    assert body.text == "Revenue grew 12% in Q3."
    assert [r.text for r in body.runs if r.bold] == ["12%"]
    assert [r.text for r in body.runs if r.italic] == ["Q3"]

    assert len(document.tables) == 1
    table = document.tables[0]
    assert [c.text for c in table.rows[0].cells] == ["Region", "Revenue"]
    assert [c.text for c in table.rows[2].cells] == ["APAC", "12"]

    page_breaks = document.element.body.xpath('.//w:br[@w:type="page"]')
    assert len(page_breaks) == 1


def test_control_characters_are_removed() -> None:
    spec = _spec(blocks=[{"type": "paragraph", "text": "bad\x00\x0bchars"}])
    document = Document(BytesIO(build_docx(spec)))
    assert any(p.text == "badchars" for p in document.paragraphs)


def test_rejects_too_many_blocks() -> None:
    blocks = [{"type": "page_break"}] * (MAX_DOCUMENT_BLOCKS + 1)
    with pytest.raises(ValidationError):
        _spec(blocks=blocks)


@pytest.mark.parametrize(
    "block",
    [
        {"type": "heading", "text": "H", "level": 4},
        {"type": "table", "header": ["a", "b"], "rows": [["only one"]]},
        {"type": "bullet_list", "items": []},
        {"type": "unknown", "text": "x"},
    ],
)
def test_rejects_invalid_blocks(block: dict[str, Any]) -> None:
    with pytest.raises(ValidationError):
        _spec(blocks=[block])


def test_rejects_empty_blocks() -> None:
    with pytest.raises(ValidationError):
        _spec(blocks=[])


@pytest.mark.parametrize(
    ("raw", "extension", "expected"),
    [
        ("Report", ".docx", "Report.docx"),
        ("report.pdf", ".docx", "report.docx"),
        ("../../etc/passwd", ".docx", "passwd.docx"),
        ("C:\\tmp\\plan.xlsx", ".xlsx", "plan.xlsx"),
        ('a<b>:"c"|?*', ".pptx", "a_b_c_.pptx"),
        ("...", ".docx", "file.docx"),
        ("  spaced   name  ", ".docx", "spaced name.docx"),
    ],
)
def test_sanitize_file_name(raw: str, extension: str, expected: str) -> None:
    assert sanitize_file_name(raw, extension) == expected


def test_sanitize_file_name_bounds_length() -> None:
    name = sanitize_file_name("x" * 500, ".docx")
    assert name.endswith(".docx")
    assert len(name) <= 120


def test_parse_inline_runs() -> None:
    runs = parse_inline_runs("a **b** c *d* e")
    assert [(r.text, r.bold, r.italic) for r in runs] == [
        ("a ", False, False),
        ("b", True, False),
        (" c ", False, False),
        ("d", False, True),
        (" e", False, False),
    ]
