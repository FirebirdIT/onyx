from io import BytesIO
from typing import Any

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from onyx.tools.tool_implementations.office.xlsx_builder import (
    MAX_ROWS_PER_SHEET,
    WorkbookSpec,
    build_xlsx,
)


def _sheet(**overrides: Any) -> dict[str, Any]:
    sheet: dict[str, Any] = {
        "name": "Sales",
        "header": ["Region", "Q1", "Q2", "Active"],
        "rows": [
            ["EMEA", 10, 12.5, True],
            ["APAC", 8, 9, False],
            ["Total", "=SUM(B2:B3)", "=SUM(C2:C3)", None],
        ],
        "column_widths": [18, 10],
        "charts": [
            {
                "type": "bar",
                "title": "Revenue",
                "data_range": "B1:C3",
                "categories_range": "A2:A3",
            },
            {"type": "pie", "data_range": "B1:B3", "anchor": "H20"},
        ],
    }
    sheet.update(overrides)
    return sheet


def _spec(sheets: list[dict[str, Any]]) -> WorkbookSpec:
    return WorkbookSpec.model_validate({"file_name": "sales", "sheets": sheets})


def test_build_xlsx_round_trip() -> None:
    spec = _spec([_sheet(), _sheet(name="Notes", charts=[], freeze_header=False)])
    workbook = load_workbook(BytesIO(build_xlsx(spec)))

    assert workbook.sheetnames == ["Sales", "Notes"]
    sheet = workbook["Sales"]

    def value(row: int, column: int) -> object:
        return sheet.cell(row=row, column=column).value

    assert [value(1, column) for column in range(1, 5)] == [
        "Region",
        "Q1",
        "Q2",
        "Active",
    ]
    assert value(2, 1) == "EMEA"
    assert value(2, 2) == 10
    assert value(2, 3) == 12.5
    assert value(2, 4) is True
    assert value(3, 4) is False
    assert value(4, 2) == "=SUM(B2:B3)"
    assert sheet.cell(row=4, column=2).data_type == "f"
    assert sheet.cell(row=1, column=1).font.b is True
    assert sheet.freeze_panes == "A2"
    # The openpyxl stubs omit these Worksheet attributes.
    dimensions = sheet.column_dimensions  # ty: ignore[unresolved-attribute]
    assert dimensions["A"].width == 18
    assert len(sheet._charts) == 2  # ty: ignore[unresolved-attribute]

    notes = workbook["Notes"]
    assert notes.freeze_panes is None
    assert len(notes._charts) == 0  # ty: ignore[unresolved-attribute]


@pytest.mark.parametrize(
    "sheet",
    [
        _sheet(name="bad/name"),
        _sheet(name="x" * 32),
        _sheet(rows=[["a", 1, 2, True, "extra"]]),
        _sheet(rows=[["=cmd|' /C calc'!A0"]]),
        _sheet(rows=[[float("nan")]]),
        _sheet(charts=[{"type": "bar", "data_range": "B1:Z99"}]),
        _sheet(charts=[{"type": "pie", "data_range": "B1:C3"}]),
        _sheet(charts=[{"type": "line", "data_range": "not a range"}]),
        _sheet(charts=[{"type": "area", "data_range": "B1:B3"}]),
        _sheet(column_widths=[10, 10, 10, 10, 10]),
    ],
)
def test_rejects_invalid_sheet(sheet: dict[str, Any]) -> None:
    with pytest.raises(ValidationError):
        _spec([sheet])


def test_rejects_duplicate_sheet_names() -> None:
    with pytest.raises(ValidationError):
        _spec([_sheet(name="Data"), _sheet(name="data")])


def test_rejects_too_many_rows() -> None:
    rows = [["r", 1, 2, True]] * (MAX_ROWS_PER_SHEET + 1)
    with pytest.raises(ValidationError):
        _spec([_sheet(rows=rows, charts=[])])


def test_rejects_no_sheets() -> None:
    with pytest.raises(ValidationError):
        _spec([])
